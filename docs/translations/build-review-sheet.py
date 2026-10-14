#!/usr/bin/env python3
"""Build docs/translations/review.xlsx from lang/{ms,id}.json for native review.

Usage (repo root): python3 docs/translations/build-review-sheet.py
Needs: openpyxl (pip install openpyxl) and php on PATH (reads the passkeys lang files).

Sheets:
  App strings        one row per app string: where used, English, Malay, Indonesian, notes.
                     The scan mirrors the static guard in
                     tests/Feature/Localisation/TranslationCompletenessTest.php, so the row
                     count equals the number of keys that guard finds.
  Passkey strings    overrides for spatie/laravel-passkeys (lang/vendor/passkeys/{ms,id}).
  Framework strings  laravel-lang entries whose keys are not in our code.
Privacy-policy rows are highlighted for legal review. Output is deterministic.
"""
import datetime, json, pathlib, re, subprocess, zipfile
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

ROOT = pathlib.Path(__file__).resolve().parents[2]
SCAN_PATHS = ["app", "resources/views"]
# Keep in sync with keysInCode() in TranslationCompletenessTest.php.
PATTERN = re.compile(r"""(?<![\w>:$])(?:__|trans_choice|trans|@lang)\(\s*(['"])((?:\\.|(?!\1).)*)\1""", re.S)


def app_key(k):
    # Mirrors AppTranslationKey::isAppKey().
    return "::" not in k and not re.fullmatch(r"[a-z_]+\.[a-z0-9_.]+", k)


def unescape(raw):
    # Equivalent of PHP stripcslashes() for the sequences that occur in literals.
    return bytes(raw, "utf-8").decode("unicode_escape").encode("latin-1").decode("utf-8")


def php_array(path):
    out = subprocess.run(["php", "-r", "echo json_encode(require $argv[1], JSON_UNESCAPED_UNICODE);", str(path)],
                         check=True, capture_output=True, text=True).stdout
    return json.loads(out)


uses = {}
files = sorted(p for d in SCAN_PATHS for p in (ROOT / d).rglob("*.php"))
for path in files:
    src = path.read_text(encoding="utf-8")
    for m in PATTERN.finditer(src):
        key = unescape(m.group(2))
        if app_key(key):
            line = src.count("\n", 0, m.start()) + 1
            uses.setdefault(key, []).append(f"{path.relative_to(ROOT)}:{line}")

ms = json.loads((ROOT / "lang/ms.json").read_text(encoding="utf-8"))
idn = json.loads((ROOT / "lang/id.json").read_text(encoding="utf-8"))

wb = Workbook()
wb.properties.creator = "build-review-sheet.py"
wb.properties.created = wb.properties.modified = datetime.datetime(2000, 1, 1)

ws = wb.active
ws.title = "App strings"
ws.append(["Where used", "English", "Bahasa Melayu", "Bahasa Indonesia", "Reviewer notes"])
legal = PatternFill("solid", fgColor="FFF2CC")
for key in sorted(uses, key=lambda k: (uses[k][0], k)):
    ws.append(["\n".join(uses[key]), key, ms.get(key, ""), idn.get(key, ""), ""])
    if any("privacy.blade.php" in u for u in uses[key]):
        for cell in ws[ws.max_row]:
            cell.fill = legal
        ws.cell(ws.max_row, 5, "LEGAL: privacy policy — needs accountable review")

pk = wb.create_sheet("Passkey strings")
pk.append(["English", "Bahasa Melayu", "Bahasa Indonesia", "Reviewer notes"])
pk_en = php_array(ROOT / "vendor/spatie/laravel-passkeys/resources/lang/en/passkeys.php")
pk_ms = php_array(ROOT / "lang/vendor/passkeys/ms/passkeys.php")
pk_id = php_array(ROOT / "lang/vendor/passkeys/id/passkeys.php")
for key in sorted(set(pk_en) | set(pk_ms) | set(pk_id)):
    pk.append([pk_en.get(key, ""), pk_ms.get(key, ""), pk_id.get(key, ""), f"package key: passkeys.{key}"])

fw = wb.create_sheet("Framework strings")
fw.append(["English", "Bahasa Melayu", "Bahasa Indonesia"])
for key in sorted(set(ms) | set(idn)):
    if key not in uses:
        fw.append([key, ms.get(key, ""), idn.get(key, "")])

for sheet, widths in ((ws, [38, 48, 48, 48, 30]), (pk, [48, 48, 48, 30]), (fw, [48, 48, 48])):
    sheet.freeze_panes = "A2"
    for i, w in enumerate(widths):
        sheet.column_dimensions[chr(65 + i)].width = w
    for cell in sheet[1]:
        cell.font = Font(bold=True)
    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")

out = ROOT / "docs/translations/review.xlsx"
wb.save(out)

# openpyxl stamps zip entries and dcterms:modified with the current time; pin both so reruns are byte-identical.
with zipfile.ZipFile(out) as zin:
    entries = [(i.filename, zin.read(i)) for i in zin.infolist()]
with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as zout:
    for name, data in entries:
        if name == "docProps/core.xml":
            data = re.sub(rb"(<dcterms:modified[^>]*>)[^<]*", rb"\g<1>2000-01-01T00:00:00Z", data)
        info = zipfile.ZipInfo(name, date_time=(2000, 1, 1, 0, 0, 0))
        info.compress_type = zipfile.ZIP_DEFLATED
        zout.writestr(info, data)

print(f"wrote {out.relative_to(ROOT)}: {ws.max_row - 1} app strings, "
      f"{pk.max_row - 1} passkey strings, {fw.max_row - 1} framework strings")
